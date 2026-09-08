"""Assemble the benchmarking workbook: Kalshi vs. every benchmark, per event.

Inputs (raw downloads kept untouched in data/external/):
  - Cleveland Fed daily CPI nowcasts (cleveland_nowcast_daily.csv, via pull_cleveland_nowcast.py)
  - BLS CPI-U seasonally adjusted index, current vintage (bls_cpi_cusr0000sa0.json, via pull_bls_cpi.py)
  - Our clean tables (events_table, forecast_snapshots)

Outputs (results/):
  comparison.csv, true_cpi.csv   the two evaluation tables as plain CSV
  kxcpi_benchmarks.xlsx          the same tables as a workbook for teammates:
  README              sources and caveats
  true_cpi            first-release actual (settlement) vs current BLS vintage
  comparison          one row per settled event x horizon: every forecaster side by side
  cleveland_daily     full daily nowcast archive
  consensus_template  hand-fill sheet for the economist consensus benchmark
"""
import json
import pathlib

import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
EXT = ROOT / "data" / "external"
RESULTS = ROOT / "results"
OUT = RESULTS / "kxcpi_benchmarks.xlsx"

MONTHS = {m: i + 1 for i, m in enumerate(
    ["JAN", "FEB", "MAR", "APR", "MAY", "JUN",
     "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"])}


def event_to_month(event_ticker):
    """CPI-22JUL -> 2022-07 ; KXCPI-26MAR -> 2026-03."""
    tail = event_ticker.split("-")[-1]
    return f"20{tail[:2]}-{MONTHS[tail[2:]]:02d}"


def load_bls():
    d = json.loads((EXT / "bls_cpi_cusr0000sa0.json").read_text())
    rows = [{"month": f"{p['year']}-{p['period'][1:]}", "index": p["value"]}
            for p in d["Results"]["series"][0]["data"] if p["period"].startswith("M")]
    df = pd.DataFrame(rows).sort_values("month").reset_index(drop=True)
    # '-' marks months BLS never published (2025 government-shutdown gap)
    df["index"] = pd.to_numeric(df["index"], errors="coerce")
    df["bls_mom_current"] = (df["index"].pct_change(fill_method=None) * 100).round(1)
    # a MoM change straddling an unpublished month is not a 1-month change
    df.loc[df["index"].shift(1).isna(), "bls_mom_current"] = None
    return df


def main():
    ev = pd.read_csv(ROOT / "data" / "clean" / "events_table.csv")
    sn = pd.read_csv(ROOT / "data" / "clean" / "forecast_snapshots.csv")
    nc = pd.read_csv(EXT / "cleveland_nowcast_daily.csv")
    bls = load_bls()

    ev["month"] = ev.event_ticker.map(event_to_month)
    ev["release_date"] = pd.to_datetime(ev.close_time).dt.date.astype(str)
    settled = ev[ev.settled & ev.actual_cpi_mom.notna()].sort_values("month")

    # ---- true_cpi sheet: first release vs current BLS vintage ----
    true_cpi = settled[["month", "event_ticker", "release_date",
                        "actual_cpi_mom"]].rename(
        columns={"actual_cpi_mom": "first_release_mom"})
    true_cpi = true_cpi.merge(bls[["month", "index", "bls_mom_current"]],
                              on="month", how="left")
    true_cpi["revised_since_release"] = (
        (true_cpi.first_release_mom - true_cpi.bls_mom_current).abs() > 0.049
    ).map({True: "YES", False: ""})
    true_cpi = true_cpi.rename(columns={"index": "bls_index_current"})

    # released first-prints in release order, for the naive baselines
    rel = settled[["release_date", "month", "actual_cpi_mom"]].sort_values("release_date")

    # ---- comparison sheet ----
    rows = []
    for _, s in sn.sort_values(["event_ticker", "horizon"]).iterrows():
        month = event_to_month(s.event_ticker)
        snap_date = str(s.date)
        # Cleveland: latest nowcast for this target month on/before the snapshot date
        g = nc[(nc.target_month == month) & (nc.nowcast_date <= snap_date)]
        cleveland = round(g.cpi_mom_nowcast.iloc[-1], 3) if len(g) else None
        # naive: prints already released as of the snapshot date (real-time discipline)
        known = rel[rel.release_date <= snap_date].actual_cpi_mom
        naive_last = known.iloc[-1] if len(known) else None
        naive_avg12 = round(known.tail(12).mean(), 2) if len(known) >= 6 else None

        actual = s.actual_cpi_mom
        row = {
            "event_ticker": s.event_ticker, "month": month,
            "horizon_days": int(s.horizon), "snapshot_date": snap_date,
            "release_date": str(s.release_date), "actual_first_release": actual,
            "kalshi_implied_mean": round(s.implied_mean, 3),
            "cleveland_nowcast": cleveland,
            "naive_last_print": naive_last,
            "naive_avg_12m": naive_avg12,
            "consensus_survey": None,        # hand-fill via consensus_template
        }
        for name, fc in [("kalshi", row["kalshi_implied_mean"]),
                         ("cleveland", cleveland),
                         ("naive_last", naive_last),
                         ("naive_avg", naive_avg12)]:
            row[f"abs_err_{name}"] = (round(abs(fc - actual), 3)
                                      if fc is not None and pd.notna(fc) else None)
        rows.append(row)
    comp = pd.DataFrame(rows)

    # ---- consensus template ----
    cons = settled[["event_ticker", "month", "release_date"]].copy()
    cons["consensus_mom"] = None
    cons["source_url"] = None
    cons.loc[cons.index[-1], "consensus_mom"] = "e.g. 0.2"
    cons.loc[cons.index[-1], "source_url"] = "e.g. reuters.com/... (delete this example row's values)"

    readme = pd.DataFrame({"About this workbook": [
        "Benchmarks for evaluating Kalshi's KXCPI forecasts. Raw downloads live untouched in data/external/.",
        "",
        "Sources:",
        "  Cleveland Fed daily CPI nowcast - clevelandfed.org/indicators-and-data/inflation-nowcasting (chart JSON, full archive to 2013). Pulled by src/pull_cleveland_nowcast.py.",
        "  BLS CPI-U, seasonally adjusted index CUSR0000SA0 - api.bls.gov public API, CURRENT vintage (reflects later revisions).",
        "  First-release actuals - Kalshi settlement values (the number the market actually settled on).",
        "  Economist consensus - no free archive exists; hand-fill the consensus_template sheet from press coverage of each release.",
        "",
        "Sheets:",
        "  true_cpi - first-release print vs today's revised BLS figure; 'revised_since_release' flags months where they differ. EVALUATE AGAINST first_release_mom.",
        "  comparison - one row per settled event per horizon (30/14/7/1 days pre-release): all forecasters side by side, sampled at the SAME snapshot date, with absolute errors.",
        "  cleveland_daily - the full daily nowcast archive (2013-2026).",
        "  consensus_template - fill consensus_mom per month from news archives, then merge into comparison.",
        "",
        "Fairness note: Cleveland nowcast is matched to the latest value on/before each Kalshi snapshot date; naive baselines use only prints already RELEASED by that date (no look-ahead).",
        "Workbook generated by src/build_benchmarks.py.",
    ]})

    RESULTS.mkdir(exist_ok=True)
    true_cpi.to_csv(RESULTS / "true_cpi.csv", index=False)
    comp.to_csv(RESULTS / "comparison.csv", index=False)

    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        readme.to_excel(xw, sheet_name="README", index=False)
        true_cpi.to_excel(xw, sheet_name="true_cpi", index=False)
        comp.to_excel(xw, sheet_name="comparison", index=False)
        nc.to_excel(xw, sheet_name="cleveland_daily", index=False)
        cons.to_excel(xw, sheet_name="consensus_template", index=False)

        arial = Font(name="Arial", size=10)
        bold = Font(name="Arial", size=10, bold=True)
        fill = PatternFill("solid", fgColor="FFFF00")
        for name, df in [("README", readme), ("true_cpi", true_cpi),
                         ("comparison", comp), ("cleveland_daily", nc),
                         ("consensus_template", cons)]:
            ws = xw.book[name]
            for r in ws.iter_rows():
                for c in r:
                    c.font = arial
            for c in ws[1]:
                c.font = bold
            ws.freeze_panes = "A2"
            for i, col in enumerate(df.columns, 1):
                width = max(len(str(col)) + 2,
                            min(38, int(df[col].astype(str).str.len().fillna(0).quantile(0.9)) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
        ws = xw.book["consensus_template"]      # highlight the fill-in column
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=4).fill = fill
        xw.book["README"].column_dimensions["A"].width = 120

    print(f"wrote {OUT}, {RESULTS / 'comparison.csv'}, {RESULTS / 'true_cpi.csv'}")
    print(f"comparison rows: {len(comp)} | events: {comp.event_ticker.nunique()}")
    ok = comp[comp.horizon_days == 1].dropna(subset=["abs_err_kalshi", "abs_err_cleveland"])
    print(f"T-1 head-to-head (n={len(ok)}): "
          f"Kalshi MAE {ok.abs_err_kalshi.mean():.3f} vs "
          f"Cleveland MAE {ok.abs_err_cleveland.mean():.3f} | "
          f"Kalshi closer in {(ok.abs_err_kalshi < ok.abs_err_cleveland).sum()}, "
          f"Cleveland closer in {(ok.abs_err_cleveland < ok.abs_err_kalshi).sum()}, "
          f"ties {(ok.abs_err_kalshi == ok.abs_err_cleveland).sum()}")


if __name__ == "__main__":
    main()
